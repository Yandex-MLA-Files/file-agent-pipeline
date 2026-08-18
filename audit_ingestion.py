"""Parse and chunk a folder of documents and report anything that looks wrong.

Structural invariants, not opinions — a block with no text, a chunk over the
embedding budget, a chunk that is only a breadcrumb, a table piece that lost
its header, a duplicated chunk, a lost page number, mojibake — plus an
end-to-end check that sentences of the source survive into the chunks.

    python audit_ingestion.py ~/testdocs
    python audit_ingestion.py ~/testdocs --coverage 40

Written for a real corpus: it found the wide-table header loss and the
contentless heading chunks that §2.3 of docs/parsing_and_chunking.md describes.
"""

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent / "src"))

from file_agent.chunking import chunk_document, get_embedding_tokenizer  # noqa: E402
from file_agent.document import BlockType  # noqa: E402
from file_agent.pipeline import parse_file  # noqa: E402

SUPPORTED = (".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".md", ".txt", ".html", ".htm")
# Cyrillic text decoded with the wrong codepage.
MOJIBAKE = re.compile(r"[ÐÑ][\x80-\xbf]|[Ã][\x80-\xbf]{2}")
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
TABLE_SEPARATOR = re.compile(r"\|\s*-+\s*\|")
TIMESTAMP = re.compile(r"\[\d{1,2}:\d{2}(?::\d{2})?\]")
MIN_SENTENCE_WORDS = 6


class Findings:
    def __init__(self) -> None:
        self.counts: Counter[str] = Counter()

    def flag(self, rule: str, detail: str = "") -> None:
        self.counts[rule] += 1
        if detail and self.counts[rule] <= 3:
            print(f"    ! {rule}: {detail}")


def audit_blocks(document, findings: Findings) -> None:
    previous_key = None
    for index, block in enumerate(document.blocks):
        text = block.text or ""
        if not text.strip() and block.block_type not in (BlockType.FIGURE, BlockType.IMAGE):
            findings.flag("empty_block", f"#{index} {block.block_type.value}")
        if MOJIBAKE.search(text):
            findings.flag("mojibake", text[:60])
        if CONTROL.search(text):
            findings.flag("control_chars", repr(text[:60]))
        if document.file_type == "pdf" and block.page_number is None:
            findings.flag("pdf_block_without_page", f"#{index} {block.block_type.value}")
        if block.block_type == BlockType.HEADING and len(text) > 300:
            findings.flag("heading_too_long", text[:70])
        if not block.metadata.get("source_file"):
            findings.flag("block_without_source_file", f"#{index}")
        key = (block.block_type, text.strip())
        if text.strip() and key == previous_key:
            findings.flag("duplicate_adjacent_block", text[:60])
        previous_key = key


def audit_chunks(chunks, budget: int, findings: Findings) -> None:
    tokenizer = get_embedding_tokenizer()
    seen: Counter[str] = Counter()
    for chunk in chunks:
        text = chunk.text or ""
        meta = chunk.metadata or {}
        if not text.strip():
            findings.flag("empty_chunk", chunk.id)
            continue
        seen[text.strip()] += 1

        size = len(tokenizer.encode(text, add_special_tokens=False)) if tokenizer else len(text)
        if size > budget:
            findings.flag("chunk_over_budget", f"{size} > {budget}")

        crumb = " > ".join(meta.get("heading_path") or [])
        body = text[len(crumb) :].strip(" >\n") if crumb and text.startswith(crumb) else text
        if len(body) < 15:
            findings.flag("chunk_without_content", text[:70])
        if meta.get("block_type") == "table" and "|" in text and not TABLE_SEPARATOR.search(text):
            findings.flag("table_chunk_without_header", text[:70])
        if not meta.get("source_file"):
            findings.flag("chunk_without_source_file", chunk.id)

    for text, count in seen.items():
        if count > 1:
            findings.flag("duplicate_chunk", f"x{count} {text[:60]}")


def source_sentences(path: Path) -> list[str]:
    """Read the document with an independent, dumb reader."""
    suffix = path.suffix.lower()
    raw = ""
    if suffix == ".pdf":
        import fitz

        with fitz.open(str(path)) as pdf:
            raw = "\n".join(page.get_text("text") for page in pdf)
    elif suffix == ".docx":
        from docx import Document as OpenDocx

        docx = OpenDocx(str(path))
        parts = [paragraph.text for paragraph in docx.paragraphs]
        for table in docx.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        raw = "\n".join(parts)
    elif suffix in (".txt", ".md", ".html", ".htm"):
        raw = path.read_text(encoding="utf-8", errors="replace")
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=True)
        raw = "\n".join(
            str(value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
            for value in row
            if isinstance(value, str)
        )
    elif suffix == ".pptx":
        from pptx import Presentation

        raw = "\n".join(
            shape.text
            for slide in Presentation(str(path)).slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
    # Transcript timestamps are dropped by the parser on purpose.
    raw = TIMESTAMP.sub(" ", raw)
    sentences = [part.strip() for part in re.split(r"(?<=[.!?])\s+|\n", raw)]
    return [s for s in sentences if len(s.split()) >= MIN_SENTENCE_WORDS and len(s) < 300]


def normalize(text: str) -> str:
    return re.sub(r"[^0-9a-zа-яё]+", "", text.lower())


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", type=Path)
    parser.add_argument("--max-chars", type=int, default=1000)
    parser.add_argument("--overlap", type=int, default=100)
    parser.add_argument(
        "--coverage",
        type=int,
        default=0,
        metavar="N",
        help="also check that N sampled source sentences survive into the chunks",
    )
    args = parser.parse_args()

    import random

    random.seed(20260818)
    findings = Findings()
    tokenizer = get_embedding_tokenizer()
    survived = checked = 0

    for path in sorted(args.folder.iterdir()):
        if path.suffix.lower() not in SUPPORTED:
            continue
        print("=" * 78)
        try:
            document = parse_file(path)
            chunks = chunk_document(
                document=document,
                max_chars=args.max_chars,
                overlap=args.overlap,
                tokenizer=tokenizer,
            )
        except Exception as exc:  # noqa: BLE001 - the audit reports, never fails
            findings.flag("parse_failed", f"{path.name}: {type(exc).__name__}: {exc}")
            print(f"{path.name}: FAILED {type(exc).__name__}: {exc}")
            continue

        kinds = Counter(block.block_type.value for block in document.blocks)
        print(f"{path.name}: {len(document.blocks)} blocks {dict(kinds)} -> {len(chunks)} chunks")
        audit_blocks(document, findings)
        audit_chunks(chunks, args.max_chars, findings)

        if args.coverage:
            sentences = source_sentences(path)
            if sentences:
                sample = random.sample(sentences, min(args.coverage, len(sentences)))
                haystack = normalize(
                    " ".join(
                        chunk.text + " " + str((chunk.metadata or {}).get("context") or "")
                        for chunk in chunks
                    )
                )
                missing = [s for s in sample if normalize(s) not in haystack]
                checked += len(sample)
                survived += len(sample) - len(missing)
                print(f"    coverage: {len(sample) - len(missing)}/{len(sample)} sentences")
                for sentence in missing[:3]:
                    print(f"    ? not found: {sentence[:100]!r}")

    print("=" * 78)
    print("FINDINGS")
    for rule, count in findings.counts.most_common():
        print(f"  {count:5d}  {rule}")
    if not findings.counts:
        print("  none")
    if checked:
        print(f"COVERAGE {survived}/{checked} sampled sentences survive parsing and chunking")
    return 0


if __name__ == "__main__":
    sys.exit(main())
