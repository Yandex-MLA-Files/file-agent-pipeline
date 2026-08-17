"""Structured PPTX parser built on python-pptx.

Every slide becomes a small section: a ``HEADING`` block with the slide title
(``hierarchy_level`` 1; slides that only carry a section-divider title become
level-1 headings for the slides that follow, whose titles then get level 2),
followed by the slide body in visual reading order (top-to-bottom,
left-to-right; grouped shapes are flattened). Bullet paragraphs are rendered as
``LIST`` blocks with their indentation, tables as Markdown ``TABLE`` blocks,
pictures as ``FIGURE`` blocks carrying the image bytes (so a VLM can describe
them), charts as their series data, and speaker notes as a trailing ``TEXT``
block. Slide numbers are stored as ``page_number`` so citations work exactly
like PDF pages.
"""

import logging
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

from file_agent.document import BlockType, Document
from file_agent.parsers.base import BaseParser
from file_agent.parsers.common import (
    BlockFactory,
    clean_text,
    format_cell,
    list_to_markdown,
    strip_bullet,
    table_to_markdown,
)
from file_agent.telemetry import tracer

logger = logging.getLogger(__name__)

# Text frames whose only content is a slide number / footer date add nothing.
_SLIDE_NUMBER = re.compile(r"^\s*\d{1,3}\s*(/\s*\d{1,3})?\s*$")
_MIN_IMAGE_BYTES = 2048
# Auto-generated shape names in any language: "Picture 3", "Рисунок 2", "图片 8",
# "Замещающая рамка рисунка 2" — a short non-numeric prefix and a counter.
_GENERIC_SHAPE_NAME = re.compile(r"^[^\d]{0,40}\d+$")
# Default presentation titles that Office writes into core properties.
_DEFAULT_DECK_TITLE = re.compile(
    r"^(powerpoint\s+(presentation|演示文稿)|презентация\s+powerpoint|presentation\d*|slide\s*1)$",
    re.IGNORECASE,
)


class PPTXParser(BaseParser):
    def parse(self, file_path: Path) -> Document:
        path = Path(file_path)
        with tracer.start_as_current_span("file_agent.pptx_parse") as span:
            span.set_attribute("file_agent.file_name", path.name)

            presentation = Presentation(str(path))
            factory = BlockFactory(source_file=path.name)
            title = _presentation_title(presentation)
            figure_count = 0
            table_count = 0
            in_section = False

            slide_height = _emu(presentation.slide_height)
            for slide_index, slide in enumerate(presentation.slides, start=1):
                slide_title = _slide_title(slide, slide_height)
                body_shapes = _ordered_shapes(slide)
                is_divider = _is_section_divider(slide, slide_title, body_shapes)

                if slide_title:
                    level = 1 if (is_divider or not in_section) else 2
                    if is_divider:
                        in_section = True
                    factory.heading(
                        slide_title,
                        level,
                        page_number=slide_index,
                        metadata={"slide_number": slide_index, "slide_layout": _layout_name(slide)},
                    )
                    if title is None:
                        title = slide_title
                elif not in_section:
                    # Untitled slide: keep it findable by number.
                    factory.heading(
                        f"Slide {slide_index}",
                        2,
                        page_number=slide_index,
                        metadata={"slide_number": slide_index, "synthetic_title": True},
                    )

                emitter = _SlideEmitter(factory, slide_index, slide_title)
                for shape in body_shapes:
                    emitter.emit(shape)
                figure_count += emitter.figure_count
                table_count += emitter.table_count

                notes = _speaker_notes(slide)
                if notes:
                    factory.add(
                        f"Speaker notes: {notes}",
                        BlockType.TEXT,
                        {"slide_number": slide_index, "speaker_notes": True},
                        page_number=slide_index,
                    )

            if not factory.blocks:
                factory.add(
                    "",
                    BlockType.TEXT,
                    {"warning": "empty presentation"},
                    page_number=1,
                    skip_empty=False,
                )

            document = Document(
                file_name=path.name,
                file_type="pptx",
                blocks=factory.blocks,
                metadata={
                    "parsing_method": "python-pptx",
                    "title": title,
                    "slide_count": len(presentation.slides),
                    "figure_count": figure_count,
                    "table_count": table_count,
                },
            )
            document.build_table_of_contents()
            span.set_attribute("file_agent.block_count", len(document.blocks))
            logger.info("Parsed %s into %d block(s)", path.name, len(document.blocks))
            return document


class _SlideEmitter:
    def __init__(self, factory: BlockFactory, slide_number: int, slide_title: str | None) -> None:
        self.factory = factory
        self.slide_number = slide_number
        self.slide_title = slide_title
        self.figure_count = 0
        self.table_count = 0
        self._meta = {"slide_number": slide_number}

    def emit(self, shape: Any) -> None:
        if getattr(shape, "has_table", False) and shape.has_table:
            self._emit_table(shape)
            return
        if getattr(shape, "has_chart", False) and shape.has_chart:
            self._emit_chart(shape)
            return
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
            self._emit_picture(shape)
            return
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            self._emit_text_frame(shape)

    def _emit_text_frame(self, shape: Any) -> None:
        paragraphs = list(shape.text_frame.paragraphs)
        text_all = clean_text("\n".join(p.text for p in paragraphs))
        if not text_all or text_all == (self.slide_title or ""):
            return
        if _SLIDE_NUMBER.match(text_all) or _is_footer_placeholder(shape):
            return

        # Bullet levels are the only structure PowerPoint gives. Explicit
        # bullets and indented paragraphs are list items; so are the paragraphs
        # of a body placeholder (its bullets are inherited from the layout and
        # invisible in the slide XML). Anything else is prose.
        body_placeholder = _is_body_placeholder(shape)
        non_empty = [p for p in paragraphs if clean_text(p.text)]
        items: list[tuple[int, str]] = []
        prose: list[str] = []
        for paragraph in non_empty:
            text = clean_text(paragraph.text)
            level = int(getattr(paragraph, "level", 0) or 0)
            bulleted = (
                level > 0 or _has_bullet(paragraph) or (body_placeholder and len(non_empty) > 1)
            )
            if bulleted:
                items.append((level, strip_bullet(text)))
            else:
                prose.append(text)

        if items and not prose:
            if len(items) == 1 and items[0][0] == 0:
                self.factory.add(items[0][1], BlockType.TEXT, self._meta, self.slide_number)
            else:
                markdown = list_to_markdown(items)
                self.factory.add(
                    markdown,
                    BlockType.LIST,
                    {**self._meta, "item_count": len(items)},
                    self.slide_number,
                )
        else:
            combined = "\n".join(prose + [f"- {text}" for _, text in items])
            self.factory.add(combined, BlockType.TEXT, self._meta, self.slide_number)

    def _emit_table(self, shape: Any) -> None:
        rows: list[list[str]] = []
        for row in shape.table.rows:
            rows.append([clean_text(cell.text).replace("\n", " ") for cell in row.cells])
        markdown = table_to_markdown(rows)
        if not markdown:
            return
        self.table_count += 1
        self.factory.add(
            markdown,
            BlockType.TABLE,
            {**self._meta, "row_count": len(rows)},
            self.slide_number,
        )

    def _emit_chart(self, shape: Any) -> None:
        try:
            chart = shape.chart
            title = chart.chart_title.text_frame.text if chart.has_title else ""
            rows: list[list[Any]] = []
            categories = list(chart.plots[0].categories) if chart.plots else []
            header = ["category"] + [series.name for plot in chart.plots for series in plot.series]
            rows.append(header)
            values_by_series = [
                list(series.values) for plot in chart.plots for series in plot.series
            ]
            for index, category in enumerate(categories):
                row = [format_cell(category)]
                for values in values_by_series:
                    row.append(values[index] if index < len(values) else "")
                rows.append(row)
            markdown = table_to_markdown(rows)
        except Exception:  # pragma: no cover - charts vary a lot; never fail the slide
            logger.debug("Could not read chart on slide %s", self.slide_number, exc_info=True)
            return
        if not markdown:
            return
        text = f"Chart: {title}\n{markdown}" if title else f"Chart:\n{markdown}"
        self.factory.add(text, BlockType.TABLE, {**self._meta, "chart": True}, self.slide_number)

    def _emit_picture(self, shape: Any) -> None:
        try:
            blob = shape.image.blob
        except Exception:
            return
        if not blob or len(blob) < _MIN_IMAGE_BYTES:
            return
        self.figure_count += 1
        caption = clean_text(getattr(shape, "name", "") or "")
        # Auto-generated shape names ("Picture 3", "Рисунок 2", "图片 8") carry no meaning.
        if _GENERIC_SHAPE_NAME.match(caption):
            caption = ""
        alt_text = _alt_text(shape)
        text = alt_text or caption
        self.factory.add(
            text,
            BlockType.FIGURE,
            {**self._meta, "figure_index": self.figure_count},
            self.slide_number,
            image_bytes=blob,
            skip_empty=False,
        )


# -- helpers --------------------------------------------------------------------


def _presentation_title(presentation: Any) -> str | None:
    core = getattr(presentation, "core_properties", None)
    title = getattr(core, "title", None)
    if not isinstance(title, str) or not title.strip():
        return None
    title = " ".join(title.split())
    return None if _DEFAULT_DECK_TITLE.match(title) else title


def _layout_name(slide: Any) -> str:
    try:
        return slide.slide_layout.name
    except Exception:  # pragma: no cover
        return ""


def _slide_title(slide: Any, slide_height: int) -> str | None:
    try:
        title_shape = slide.shapes.title
    except Exception:  # pragma: no cover
        title_shape = None
    if title_shape is not None and title_shape.has_text_frame:
        text = clean_text(title_shape.text_frame.text)
        if text:
            return " ".join(text.split())
    # Layouts without a title placeholder: the title is the short text frame
    # with the largest font in the upper part of the slide (top-most on ties).
    candidates = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
            continue
        text = clean_text(shape.text_frame.text)
        if not text or len(text) > 140 or "\n" in text or _SLIDE_NUMBER.match(text):
            continue
        if _is_footer_placeholder(shape):
            continue
        top = _emu(shape.top)
        if slide_height and top > slide_height * 0.45:
            continue
        candidates.append((-_max_font_size(shape), top, text))
    if candidates:
        candidates.sort()
        return candidates[0][2]
    return None


def _max_font_size(shape: Any) -> float:
    sizes = []
    try:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size is not None:
                    sizes.append(run.font.size.pt)
    except Exception:  # pragma: no cover
        return 0.0
    return max(sizes) if sizes else 0.0


def _emu(value: Any) -> int:
    try:
        return int(Emu(value))
    except Exception:
        return 0


def _ordered_shapes(slide: Any) -> list[Any]:
    """Flatten groups and sort shapes into reading order (rows, then columns)."""
    flat: list[Any] = []

    def visit(shape: Any) -> None:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                visit(child)
        else:
            flat.append(shape)

    for shape in slide.shapes:
        visit(shape)

    title_shape = None
    try:
        title_shape = slide.shapes.title
    except Exception:  # pragma: no cover
        pass
    body = [shape for shape in flat if shape is not title_shape]

    # Two shapes on the same "row" (tops within a band) are ordered by left.
    def key(shape: Any) -> tuple[int, int]:
        top = _emu(getattr(shape, "top", 0))
        left = _emu(getattr(shape, "left", 0))
        return (top // 300000, left)  # ~0.33 inch bands

    body.sort(key=key)
    return body


def _is_section_divider(slide: Any, slide_title: str | None, body_shapes: list[Any]) -> bool:
    if not slide_title:
        return False
    layout = _layout_name(slide).lower()
    if "section" in layout or "раздел" in layout or "title only" in layout:
        return True
    text_shapes = [
        shape
        for shape in body_shapes
        if getattr(shape, "has_text_frame", False)
        and shape.has_text_frame
        and clean_text(shape.text_frame.text)
        and not _SLIDE_NUMBER.match(clean_text(shape.text_frame.text))
    ]
    return not text_shapes and not any(
        getattr(shape, "has_table", False) or hasattr(shape, "image") for shape in body_shapes
    )


def _speaker_notes(slide: Any) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        return clean_text(slide.notes_slide.notes_text_frame.text)
    except Exception:  # pragma: no cover
        return ""


def _is_footer_placeholder(shape: Any) -> bool:
    try:
        if not shape.is_placeholder:
            return False
        kind = shape.placeholder_format.type
    except Exception:
        return False
    return kind in (
        PP_PLACEHOLDER.FOOTER,
        PP_PLACEHOLDER.SLIDE_NUMBER,
        PP_PLACEHOLDER.DATE,
    )


def _is_body_placeholder(shape: Any) -> bool:
    try:
        if not shape.is_placeholder:
            return False
        kind = shape.placeholder_format.type
    except Exception:
        return False
    return kind in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE)


def _has_bullet(paragraph: Any) -> bool:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return False
    return any(child.tag.endswith(("buChar", "buAutoNum")) for child in p_pr.iterchildren())


def _alt_text(shape: Any) -> str:
    try:
        descr = shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr")
    except Exception:
        return ""
    if descr and isinstance(descr[0], str):
        text = clean_text(descr[0])
        # Office auto-generates "Изображение выглядит как ..." / "A picture containing"
        if re.match(r"^(изображение выглядит как|a picture containing|image)", text, re.I):
            return ""
        return text
    return ""
