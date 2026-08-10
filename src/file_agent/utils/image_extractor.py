import io
import logging
from pathlib import Path

import fitz  # PyMuPDF
from PIL import Image
from pptx import Presentation

logger = logging.getLogger(__name__)


def extract_image_from_pptx(
    pptx_path: Path,
    slide_number: int,
    shape_index: int,
) -> Image.Image | None:
    """Read a picture shape's raw image bytes from a PPTX slide.

    Unlike PDF, PPTX pictures need no rendering/cropping - python-pptx exposes
    the original embedded image bytes directly.

    :param pptx_path: path to the PPTX file.
    :param slide_number: 1-indexed slide number.
    :param shape_index: position of the picture shape in ``slide.shapes``.
    :return: a PIL.Image, or None if the shape could not be read.
    """
    try:
        presentation = Presentation(str(pptx_path))
        slide = presentation.slides[slide_number - 1]
        shape = slide.shapes[shape_index]
        return Image.open(io.BytesIO(shape.image.blob))
    except Exception as exc:
        logger.warning(
            "Could not extract image from slide %s, shape %s: %s", slide_number, shape_index, exc
        )
        return None


def extract_image_from_pdf(
    pdf_path: Path,
    page_number: int,
    bbox: tuple[float, float, float, float],
    zoom: float = 2.0,
) -> Image.Image | None:
    """Render a region of a PDF page to a PIL image.

    :param pdf_path: path to the PDF file.
    :param page_number: 1-indexed page number.
    :param bbox: region to crop as (x0, y0, x1, y1) in top-left PDF coordinates.
    :param zoom: scale factor to improve resolution of the crop (default 2.0).
    :return: a PIL.Image, or None if the region could not be rendered.
    """
    try:
        with fitz.open(str(pdf_path)) as doc:
            page = doc[page_number - 1]
            rect = fitz.Rect(bbox)
            matrix = fitz.Matrix(zoom, zoom)
            pixmap = page.get_pixmap(matrix=matrix, clip=rect)
            return Image.open(io.BytesIO(pixmap.tobytes("png")))
    except Exception as exc:
        logger.warning("Could not extract image from page %s, bbox=%s: %s", page_number, bbox, exc)
        return None
