import io

import docx
from docx.shared import Pt
from PIL import Image

from file_agent.document import BlockType
from file_agent.parsers.docx_parser import DOCXParser
from file_agent.pipeline import parse_file


def _png_bytes(size=(200, 150)) -> bytes:
    buffer = io.BytesIO()
    image = Image.new("RGB", size, "white")
    # A gradient keeps the PNG above the "decorative 1-pixel image" size floor.
    for x in range(size[0]):
        for y in range(0, size[1], 3):
            image.putpixel((x, y), (x % 256, y % 256, (x * y) % 256))
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def create_docx(file_path):
    document = docx.Document()
    document.add_heading("Программа экзамена", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Целью ").bold = True
    paragraph.add_run("испытания является оценка уровня освоения ")
    paragraph.add_run("компетенций.").italic = True
    document.add_heading("1.2 Раздел второй", level=2)
    document.add_paragraph("Первый пункт", style="List Bullet")
    document.add_paragraph("Второй пункт", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Метрика"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "Точность"
    table.cell(1, 1).text = "0.9"
    document.add_picture(io.BytesIO(_png_bytes()))
    document.add_paragraph("Рисунок 1 – Схема процесса")
    code = document.add_paragraph()
    run = code.add_run("print('hi')")
    run.font.name = "Courier New"
    run.font.size = Pt(10)
    document.save(file_path)


def test_docx_parser_keeps_paragraphs_whole_and_types_blocks(tmp_path):
    file_path = tmp_path / "sample.docx"
    create_docx(file_path)

    document = DOCXParser().parse(file_path)

    assert document.file_type == "docx"
    assert document.metadata["parsing_method"] == "python-docx"
    assert document.metadata["title"] == "Программа экзамена"
    types = [b.block_type for b in document.blocks]
    assert types == [
        BlockType.HEADING,
        BlockType.TEXT,
        BlockType.HEADING,
        BlockType.LIST,
        BlockType.TABLE,
        BlockType.FIGURE,
        BlockType.CODE,
    ]
    # Runs with different formatting are one paragraph, not three fragments.
    assert document.blocks[1].text == "Целью испытания является оценка уровня освоения компетенций."
    assert document.blocks[0].metadata["hierarchy_level"] == 1
    assert document.blocks[2].metadata["hierarchy_level"] == 2
    assert document.blocks[3].text == "- Первый пункт\n- Второй пункт"
    assert document.blocks[4].text.splitlines() == [
        "| Метрика | Значение |",
        "| --- | --- |",
        "| Точность | 0.9 |",
    ]
    figure = document.blocks[5]
    assert figure.image_bytes is not None
    assert figure.metadata["caption"] == "Рисунок 1 – Схема процесса"
    assert figure.text == "Рисунок 1 – Схема процесса"
    assert document.blocks[6].text == "print('hi')"


def test_pipeline_routes_docx_to_python_docx_parser(tmp_path, monkeypatch):
    file_path = tmp_path / "sample.docx"
    create_docx(file_path)
    monkeypatch.setenv("VLM_BACKEND", "off")

    document = parse_file(file_path)

    assert document.metadata["parsing_method"] == "python-docx"
    assert document.blocks[0].block_type == BlockType.HEADING


def _rich_document(tmp_path):
    from tests._docx_fixture import write_fixture

    return DOCXParser().parse(write_fixture(tmp_path / "rich.docx"))


def test_footnotes_are_marked_in_place_and_emitted_as_their_own_block(tmp_path):
    document = _rich_document(tmp_path)
    texts = [b.text for b in document.blocks]

    # The reader sees where the note was attached ...
    assert "Основной абзац со сноской. [2]" in texts
    # ... and the note itself is indexed next to the sentence it belongs to.
    note = next(b for b in document.blocks if b.metadata.get("note_type") == "footnote")
    assert note.text == "[2] Утверждён приказом № 35 от 30.12.2025."
    assert note.metadata["note_id"] == "2"
    assert document.metadata["footnote_count"] == 1


def test_list_items_carry_the_numbers_word_computes(tmp_path):
    document = _rich_document(tmp_path)
    listing = next(b for b in document.blocks if b.block_type == BlockType.LIST)

    assert listing.text.splitlines() == [
        "1. Первый пункт",
        "  1.1. Подпункт один",
        "  1.2. Подпункт два",
        "2. Второй пункт",
    ]


def test_numbered_headings_number_themselves_and_their_sub_items(tmp_path):
    """A numbered heading advances the same counter its sub-items continue.

    Counting only list paragraphs would number the second section's items
    "1.1, 1.2" — Word shows "2.1, 2.2".
    """
    from tests._docx_fixture import write_numbered_headings

    document = DOCXParser().parse(write_numbered_headings(tmp_path / "lab.docx"))
    headings = [b.text for b in document.blocks if b.block_type == BlockType.HEADING]
    lists = [b.text.splitlines() for b in document.blocks if b.block_type == BlockType.LIST]

    assert headings == ["1. Цель и содержание", "2. Порядок выполнения"]
    assert lists == [
        ["  1.1. Изучить MongoDB", "  1.2. Установить Compass"],
        ["  2.1. Создать базу", "  2.2. Проверить запросы"],
    ]


def test_a_number_already_typed_into_the_text_is_not_repeated(tmp_path):
    from file_agent.parsers.docx_parser import _repeats_marker

    assert _repeats_marker("1.2 Установка", "1.2.")
    assert _repeats_marker("8. Унифицированная система", "8.")
    assert _repeats_marker("2)", "2)")
    # A year, a quantity or a version opening the sentence is not a marker.
    assert not _repeats_marker("1996 год стал переломным", "1.")
    assert not _repeats_marker("Установка MongoDB", "3.")


def test_text_frame_becomes_its_own_block_and_leaves_the_host_paragraph_alone(tmp_path):
    document = _rich_document(tmp_path)
    frame = next(b for b in document.blocks if b.metadata.get("text_box"))

    assert frame.text == "Важно: срок хранения — 5 лет"
    assert document.metadata["text_box_count"] == 1
    # The frame's text must not have leaked into any other block.
    assert sum("срок хранения" in b.text for b in document.blocks) == 1


def test_nested_table_is_emitted_after_its_parent(tmp_path):
    document = _rich_document(tmp_path)
    tables = [b for b in document.blocks if b.block_type == BlockType.TABLE]

    assert len(tables) == 2
    assert "Раздел" in tables[0].text and "А-1" not in tables[0].text
    assert "| А-1 | 3 года |" in tables[1].text
    assert tables[1].metadata["nested_in_table"] == tables[0].metadata["table_index"]


def test_placeholder_document_properties_do_not_become_the_title(tmp_path):
    document = _rich_document(tmp_path)

    # python-docx stamps "Word Document" into core properties; the real title
    # is the first heading of the body.
    assert document.metadata["title"] == "Регламент"


def test_hyperlink_target_is_kept_next_to_its_text(tmp_path):
    """The words hide the address; a question about "where to download" needs it."""
    import docx as python_docx

    path = tmp_path / "manual.docx"
    document = python_docx.Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("скачайте установщик")
    link = paragraph._p.makeelement(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink", {}
    )
    rid = document.part.relate_to(
        "https://www.mongodb.com/try/download/compass",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", rid)
    paragraph._p.replace(run._r, link)
    link.append(run._r)
    document.save(path)

    parsed = DOCXParser().parse(path)
    body = next(b for b in parsed.blocks if "установщик" in b.text)

    assert body.text == ("скачайте установщик (https://www.mongodb.com/try/download/compass)")


def test_word_equations_are_read_as_latex(tmp_path):
    """python-docx walks past m:oMath, so equations used to vanish silently."""
    from tests._docx_fixture import write_equations

    document = DOCXParser().parse(write_equations(tmp_path / "math.docx"))
    texts = [(b.block_type.value, b.text) for b in document.blocks]

    # A paragraph that is only an equation becomes a formula block ...
    assert ("formula", r"P(A|B)=\frac{P(A \cap B)}{P(B)}") in texts
    # ... an inline one stays inside its sentence, which stops the sentence
    # from arriving as "Дисперсия равна для выборки".
    assert ("text", r"Дисперсия равна $\sigma^{2}$ для выборки.") in texts
    assert document.metadata["formula_count"] == 1


def test_equations_survive_without_the_latex_converter(tmp_path, monkeypatch):
    """Without Docling's converter the symbols are still indexed, not dropped."""
    from file_agent.parsers import docx_math
    from tests._docx_fixture import write_equations

    monkeypatch.setattr(docx_math, "_converter", None)
    monkeypatch.setattr(docx_math, "_converter_loaded", True)

    document = DOCXParser().parse(write_equations(tmp_path / "math.docx"))
    texts = [b.text for b in document.blocks]

    assert "P(A|B)=P(A∩B)P(B)" in texts
    assert any("σ2" in text for text in texts)
