import io

import fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from file_agent.document import Block, BlockType, Document
from file_agent.parsers.enhancer import DocumentEnhancer
from file_agent.parsers.pptx_parser import PPTXParser
from file_agent.vlm.base import VLMClient


class StubVLMClient(VLMClient):
    def __init__(self, description="STUB: a diagram"):
        self.description = description
        self.calls = 0

    def describe_image(self, image: Image.Image, prompt: str) -> str:
        self.calls += 1
        return self.description


class ExplodingVLMClient(VLMClient):
    def describe_image(self, image: Image.Image, prompt: str) -> str:  # pragma: no cover
        raise AssertionError("VLM must not be called")


def _make_pdf_with_shape(path):
    doc = fitz.open()
    page = doc.new_page()
    page.draw_rect(fitz.Rect(50, 50, 200, 200), color=(0, 0, 1), fill=(0, 0, 1))
    doc.save(path)
    doc.close()


def test_enhancer_describes_figure_block(tmp_path):
    pdf_path = tmp_path / "diagram.pdf"
    _make_pdf_with_shape(pdf_path)

    figure = Block(
        id="f1",
        text="",
        type="figure",
        block_type=BlockType.FIGURE,
        page_number=1,
        bbox=(50.0, 50.0, 200.0, 200.0),
    )
    document = Document(file_name="diagram.pdf", file_type="pdf", blocks=[figure])

    stub = StubVLMClient()
    DocumentEnhancer(vlm_client=stub).enhance(document, pdf_path)

    assert stub.calls == 1
    assert figure.vlm_description == "STUB: a diagram"
    assert "[Image description]: STUB: a diagram" in figure.text


def test_enhancer_skips_blocks_without_bbox(tmp_path):
    pdf_path = tmp_path / "diagram.pdf"
    _make_pdf_with_shape(pdf_path)

    text_block = Block(id="t1", text="plain", type="text", block_type=BlockType.TEXT)
    figure_no_bbox = Block(
        id="f1", text="", type="figure", block_type=BlockType.FIGURE, page_number=1
    )
    document = Document(file_name="d.pdf", file_type="pdf", blocks=[text_block, figure_no_bbox])

    # ExplodingVLMClient asserts it is never invoked.
    DocumentEnhancer(vlm_client=ExplodingVLMClient()).enhance(document, pdf_path)

    assert text_block.vlm_description is None
    assert figure_no_bbox.vlm_description is None


def _make_pptx_with_picture(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    image_bytes = io.BytesIO()
    Image.new("RGB", (10, 10), color="blue").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    slide.shapes.add_picture(
        image_bytes, left=Inches(1), top=Inches(1), width=Inches(3), height=Inches(3)
    )
    presentation.save(path)


def test_enhancer_describes_pptx_picture_block(tmp_path):
    pptx_path = tmp_path / "deck.pptx"
    _make_pptx_with_picture(pptx_path)
    document = PPTXParser().parse(pptx_path)
    picture_block = next(b for b in document.blocks if b.block_type == BlockType.IMAGE)

    stub = StubVLMClient()
    DocumentEnhancer(vlm_client=stub).enhance(document, pptx_path)

    assert stub.calls == 1
    assert picture_block.vlm_description == "STUB: a diagram"
    assert "[Image description]: STUB: a diagram" in picture_block.text


def test_enhancer_ignores_unsupported_formats(tmp_path):
    docx_path = tmp_path / "doc.docx"
    docx_path.write_bytes(b"not really a docx")

    figure = Block(
        id="f1",
        text="",
        type="figure",
        block_type=BlockType.FIGURE,
        page_number=1,
        bbox=(0.0, 0.0, 10.0, 10.0),
    )
    document = Document(file_name="doc.docx", file_type="docx", blocks=[figure])

    DocumentEnhancer(vlm_client=ExplodingVLMClient()).enhance(document, docx_path)

    assert figure.vlm_description is None
