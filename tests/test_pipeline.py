import fitz
import pytest
from openpyxl import Workbook
from pptx import Presentation

from file_agent.document import BlockType
from file_agent.pipeline import parse_file


def create_pdf(file_path, text):
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(file_path)
    document.close()


def create_xlsx(file_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Age"])
    sheet.append(["Alice", 20])
    workbook.save(file_path)
    workbook.close()


def create_pptx(file_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Project Overview"
    slide.placeholders[1].text = "PowerPoint content"
    presentation.save(file_path)


def test_parse_file_uses_markdown_parser(tmp_path):
    file_path = tmp_path / "example.md"
    file_path.write_text("Hello from pipeline", encoding="utf-8")

    document = parse_file(file_path)

    assert document.file_name == "example.md"
    assert document.file_type == "md"
    assert document.blocks[0].text == "Hello from pipeline"


def test_parse_file_uses_txt_parser(tmp_path):
    file_path = tmp_path / "example.txt"
    file_path.write_text("Hello from TXT", encoding="utf-8")

    document = parse_file(file_path)

    assert document.file_name == "example.txt"
    assert document.file_type == "txt"
    assert document.blocks[0].text == "Hello from TXT"


def test_parse_file_parses_pdf(tmp_path):
    file_path = tmp_path / "example.pdf"
    create_pdf(file_path, "Hello from PDF")

    # PDFs go through Docling by default; if Docling is unavailable the pipeline
    # gracefully falls back to the PyMuPDF parser. Both paths must recover the text.
    document = parse_file(file_path)

    assert document.file_name == "example.pdf"
    assert document.file_type == "pdf"
    assert document.blocks
    assert any("Hello from PDF" in block.text for block in document.blocks)


def test_parse_file_uses_html_parser(tmp_path):
    file_path = tmp_path / "example.html"
    file_path.write_text("<p>Hello from HTML</p>", encoding="utf-8")

    document = parse_file(file_path)

    assert document.file_name == "example.html"
    assert document.file_type == "html"
    assert len(document.blocks) == 1
    assert document.blocks[0].block_type == BlockType.TEXT
    assert document.blocks[0].text == "Hello from HTML"
    assert document.metadata["parser_profile"] == "structured"


def test_parse_file_uses_xlsx_parser(tmp_path):
    file_path = tmp_path / "example.xlsx"
    create_xlsx(file_path)

    document = parse_file(file_path)

    assert document.file_name == "example.xlsx"
    assert document.file_type == "xlsx"
    # A workbook overview opens the document, then the sheet and its table.
    assert [b.block_type for b in document.blocks] == [
        BlockType.TEXT,
        BlockType.HEADING,
        BlockType.TABLE,
    ]
    assert document.blocks[0].metadata["workbook_overview"] is True
    assert "листов — 1" in document.blocks[0].text
    assert "| Alice | 20 |" in document.blocks[2].text


def test_parse_file_uses_pptx_parser(tmp_path):
    file_path = tmp_path / "example.pptx"
    create_pptx(file_path)

    document = parse_file(file_path)

    assert document.file_name == "example.pptx"
    assert document.file_type == "pptx"
    assert document.blocks[0].block_type == BlockType.HEADING
    assert document.blocks[0].text == "Project Overview"
    assert any("PowerPoint content" in b.text for b in document.blocks)


def test_parse_file_legacy_profile_uses_flat_parsers(tmp_path, monkeypatch):
    file_path = tmp_path / "example.pptx"
    create_pptx(file_path)
    monkeypatch.setenv("PARSER_PROFILE", "legacy")

    document = parse_file(file_path)

    assert len(document.blocks) == 1
    assert document.blocks[0].type == "pptx_slide"
    assert document.metadata["parser_profile"] == "legacy"


def test_parse_file_rejects_unknown_profile(tmp_path):
    file_path = tmp_path / "example.md"
    file_path.write_text("# hi", encoding="utf-8")

    with pytest.raises(ValueError, match="PARSER_PROFILE"):
        parse_file(file_path, parser_profile="fancy")


def test_parse_file_rejects_unsupported_extension(tmp_path):
    file_path = tmp_path / "example.csv"
    file_path.write_text("Unsupported", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_file(file_path)


def test_ocr_engine_defaults_to_auto_and_validates(monkeypatch):
    from file_agent.pipeline import resolve_ocr_engine

    monkeypatch.delenv("OCR_ENGINE", raising=False)
    assert resolve_ocr_engine() == "auto"

    monkeypatch.setenv("OCR_ENGINE", "EasyOCR")
    assert resolve_ocr_engine() == "easyocr"

    monkeypatch.setenv("OCR_ENGINE", "tesseract")
    with pytest.raises(ValueError, match="OCR_ENGINE"):
        resolve_ocr_engine()


def test_auto_engine_arms_the_local_fallback_and_vlm_does_not(monkeypatch):
    from file_agent import pipeline

    monkeypatch.setattr(pipeline, "create_vlm_client", lambda: object())

    monkeypatch.setenv("OCR_ENGINE", "auto")
    assert pipeline._vlm_ocr_client().fallback is not None

    monkeypatch.setenv("OCR_ENGINE", "vlm")
    assert pipeline._vlm_ocr_client().fallback is None

    monkeypatch.setenv("OCR_ENGINE", "easyocr")
    assert pipeline._vlm_ocr_client() is None


def test_docling_parser_is_reused_between_documents_of_a_run(monkeypatch):
    """One converter per configuration: rebuilding it costs seconds per file."""
    from file_agent import pipeline

    pipeline._DOCLING_PARSERS.clear()
    built = []

    class _Parser:
        resolved_pdf_backend = None
        enrichment_available = True

        def __init__(
            self, do_ocr=False, ocr_full_page=False, enrich=None, keep_empty_regions=False
        ):
            built.append((do_ocr, ocr_full_page))

    monkeypatch.setattr(pipeline, "DoclingParser", _Parser)

    first = pipeline._docling_parser(False, False)
    assert pipeline._docling_parser(False, False) is first
    assert pipeline._docling_parser(True, False) is not first
    assert built == [(False, False), (True, False)]

    # A changed setting must not be served from the cache.
    monkeypatch.setenv("PDF_ENRICHMENT", "off")
    assert pipeline._docling_parser(False, False) is not first
    pipeline._DOCLING_PARSERS.clear()
