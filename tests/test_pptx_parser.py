from pptx import Presentation
from pptx.util import Inches

from file_agent.document import BlockType
from file_agent.parsers.legacy import PPTXParser as LegacyPPTXParser
from file_agent.parsers.pptx_parser import PPTXParser


def create_presentation(file_path):
    presentation = Presentation()

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    title_slide.shapes.title.text = "Project Overview"
    title_slide.placeholders[1].text = "This project is about document processing."

    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "Metrics"
    table = table_slide.shapes.add_table(
        rows=2,
        cols=2,
        left=Inches(1),
        top=Inches(1.5),
        width=Inches(4),
        height=Inches(1),
    ).table
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "0.85"
    table_slide.notes_slide.notes_text_frame.text = "Remember to mention the baseline."

    presentation.save(file_path)


def test_pptx_parser_emits_slide_headings_with_slide_numbers(tmp_path):
    file_path = tmp_path / "sample.pptx"
    create_presentation(file_path)

    document = PPTXParser().parse(file_path)

    assert document.file_name == "sample.pptx"
    assert document.file_type == "pptx"
    headings = [b for b in document.blocks if b.block_type == BlockType.HEADING]
    assert [h.text for h in headings] == ["Project Overview", "Metrics"]
    assert [h.page_number for h in headings] == [1, 2]
    assert document.metadata["title"] == "Project Overview"
    assert document.metadata["slide_count"] == 2
    assert [entry["title"] for entry in document.metadata["table_of_contents"]] == [
        "Project Overview",
        "Metrics",
    ]


def test_pptx_parser_extracts_body_tables_and_notes(tmp_path):
    file_path = tmp_path / "sample.pptx"
    create_presentation(file_path)

    document = PPTXParser().parse(file_path)

    texts = [b for b in document.blocks if b.block_type == BlockType.TEXT]
    assert any("This project is about document processing." == b.text for b in texts)
    tables = [b for b in document.blocks if b.block_type == BlockType.TABLE]
    assert len(tables) == 1
    assert tables[0].text.splitlines()[0] == "| Name | Value |"
    assert "| Accuracy | 0.85 |" in tables[0].text
    assert tables[0].page_number == 2
    notes = [b for b in texts if b.metadata.get("speaker_notes")]
    assert notes and "Remember to mention the baseline." in notes[0].text


def test_pptx_parser_stamps_source_and_slide_metadata(tmp_path):
    file_path = tmp_path / "sample.pptx"
    create_presentation(file_path)

    document = PPTXParser().parse(file_path)

    for block in document.blocks:
        assert block.metadata["source_file"] == "sample.pptx"
        assert block.metadata["slide_number"] == block.page_number


def test_legacy_pptx_parser_still_returns_one_block_per_slide(tmp_path):
    file_path = tmp_path / "sample.pptx"
    create_presentation(file_path)

    document = LegacyPPTXParser().parse(file_path)

    assert [b.type for b in document.blocks] == ["pptx_slide", "pptx_slide"]
    assert "Title: Project Overview" in document.blocks[0].text
    assert "Name\tValue" in document.blocks[1].text
