import datetime as dt

from openpyxl import Workbook

from file_agent.document import BlockType
from file_agent.parsers.legacy import XLSXParser as LegacyXLSXParser
from file_agent.parsers.xlsx_parser import XLSXParser


def create_workbook(file_path):
    workbook = Workbook()
    first_sheet = workbook.active
    first_sheet.title = "People"
    first_sheet.append(["Name", "Age", "City"])
    first_sheet.append(["Alice", 20, "Moscow"])
    first_sheet.append(["Bob", 22.0, "Kazan"])

    second_sheet = workbook.create_sheet("Summary")
    second_sheet.append(["Metric", "Value"])
    second_sheet.append(["Rows", 2])
    second_sheet.append([])
    second_sheet.append(["Report generated for the quarterly review of the team"])
    second_sheet.append([])
    second_sheet.append(["Date", "Amount"])
    second_sheet.append([dt.datetime(2026, 3, 31), 1100.0])

    workbook.save(file_path)
    workbook.close()


def test_xlsx_parser_emits_sheet_headings_and_markdown_tables(tmp_path):
    file_path = tmp_path / "sample.xlsx"
    create_workbook(file_path)

    document = XLSXParser().parse(file_path)

    assert document.file_name == "sample.xlsx"
    assert document.file_type == "xlsx"
    headings = [b for b in document.blocks if b.block_type == BlockType.HEADING]
    assert [h.text for h in headings] == ["Sheet: People", "Sheet: Summary"]
    tables = [b for b in document.blocks if b.block_type == BlockType.TABLE]
    assert tables[0].text.splitlines()[:3] == [
        "| Name | Age | City |",
        "| --- | --- | --- |",
        "| Alice | 20 | Moscow |",
    ]
    # 22.0 is rendered as 22, floats never carry a spurious ".0"
    assert "| Bob | 22 | Kazan |" in tables[0].text
    assert tables[0].page_number == 1
    assert tables[0].metadata["sheet_name"] == "People"


def test_xlsx_parser_splits_regions_and_keeps_notes(tmp_path):
    file_path = tmp_path / "sample.xlsx"
    create_workbook(file_path)

    document = XLSXParser().parse(file_path)

    summary = [b for b in document.blocks if b.metadata.get("sheet_name") == "Summary"]
    types = [b.block_type for b in summary]
    assert types == [BlockType.HEADING, BlockType.TABLE, BlockType.TEXT, BlockType.TABLE]
    assert "Report generated for the quarterly review" in summary[2].text
    assert "| 2026-03-31 | 1100 |" in summary[3].text
    assert document.metadata["table_count"] == 3


def test_legacy_xlsx_parser_still_returns_one_block_per_sheet(tmp_path):
    file_path = tmp_path / "sample.xlsx"
    create_workbook(file_path)

    document = LegacyXLSXParser().parse(file_path)

    assert [b.type for b in document.blocks] == ["xlsx_sheet", "xlsx_sheet"]
    assert "Name\tAge\tCity" in document.blocks[0].text
    assert document.blocks[0].metadata["sheet_name"] == "People"


def test_xlsx_parser_adds_table_profile_with_aggregates(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["id", "date", "product", "amount", "region"])
    rows = [
        (1, "2026-01-01", "Bike", 100.5, "Utah"),
        (2, "2026-01-02", "Shoes", 50.0, "Texas"),
        (3, "2026-01-03", "Bike", 200.0, "Utah"),
        (4, "2026-01-04", "Hat", 10.0, "Ohio"),
    ]
    for row in rows:
        sheet.append(list(row))
    file_path = tmp_path / "sales.xlsx"
    workbook.save(file_path)
    workbook.close()

    document = XLSXParser().parse(file_path)

    profiles = [b for b in document.blocks if b.metadata.get("table_profile")]
    assert len(profiles) == 1
    text = profiles[0].text
    assert "строк: 4" in text
    assert "amount: минимум 10 (Hat), максимум 200 (Bike), сумма 360.50" in text
    assert "region: уникальных значений 3: Utah (2), Ohio (1), Texas (1)" in text
    assert "сумма amount по region (по убыванию): Utah: 300.50, Texas: 50, Ohio: 10" in text
    # id is an identifier, not a measure
    assert "id: минимум" not in text


def test_profile_skips_prose_tables():
    """A term/definition table has no aggregates; a summary would repeat it."""
    from file_agent.parsers.table_profile import profile_table

    header = ["Термин", "Определение"]
    body = [[f"термин {i}", "очень длинное определение " * 5] for i in range(6)]

    assert profile_table(header, body) == ""


def test_profile_reports_extremes_with_their_row_label():
    from file_agent.parsers.table_profile import profile_table

    header = ["Регион", "Выручка"]
    body = [["Юг", "100"], ["Север", "900"], ["Запад", "500"], ["Восток", "300"]]

    profile = profile_table(header, body)

    assert "максимум 900 (Север)" in profile
    assert "минимум 100 (Юг)" in profile
    assert "сумма 1 800" in profile
